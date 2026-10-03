"""Task revisions R0 - R10, Approval Codes per Project, the PEC DCI Report and its exports.

Run: bench --site <site> run-tests --module pew_customizations.tests.test_task_revisions
Everything is rolled back after the run; nothing is written to disk.
"""

# ruff: noqa: RUF001 - the Latest Status texts use an en dash, as the client wrote them

import json
from io import BytesIO
from unittest.mock import patch

import frappe
from frappe.desk.query_report import export_query
from frappe.desk.query_report import run as run_report
from frappe.tests import IntegrationTestCase
from frappe.utils import add_days, getdate, today
from openpyxl import load_workbook

from pew_customizations.api.dci_export import build_dci_workbook, download_dci_xlsx
from pew_customizations.crm.opportunity import create_project
from pew_customizations.patches import migrate_dci_revisions_to_tasks as dci_patch
from pew_customizations.pew_customizations.report.pec_dci_report.pec_dci_report import execute, get_dci_data
from pew_customizations.projects.approval_codes import get_approval_codes
from pew_customizations.projects.task_revisions import refresh_latest_fields

REPORT = "PEC DCI Report"
MANAGER = "pew.test.dci.manager@example.com"
ENGINEER = "pew.test.dci.engineer@example.com"
DRAFTSMAN = "pew.test.dci.draftsman@example.com"
PEC_ADMIN = "pew.test.dci.admin@example.com"
# what the report's JS sends, so the Link filters are permission-checked as in the browser
JS_FILTERS = [
	{"fieldname": "project", "fieldtype": "Link", "options": "Project"},
	{"fieldname": "scope", "fieldtype": "Link", "options": "Scope"},
	{"fieldname": "engineer", "fieldtype": "Link", "options": "User"},
]


def _user(email, first_name, roles):
	if not frappe.db.exists("User", email):
		frappe.get_doc(
			{
				"doctype": "User",
				"email": email,
				"first_name": first_name,
				"send_welcome_email": 0,
				"roles": [{"role": r} for r in roles],
			}
		).insert(ignore_permissions=True)
	return email


class TestTaskRevisions(IntegrationTestCase):
	SHOW_TRANSACTION_COMMIT_WARNINGS = True  # everything here must be rolled back

	@classmethod
	def setUpClass(cls):
		super().setUpClass()
		cls.enterClassContext(patch("frappe.sendmail"))
		cls.company = (
			frappe.defaults.get_global_default("company") or frappe.get_all("Company", pluck="name")[0]
		)
		# the Execution Team's roles come on top of Projects User, which opens the Project itself
		_user(MANAGER, "Dci Manager", ["Projects Manager"])
		_user(ENGINEER, "Dci Engineer", ["Projects User", "Engineer"])
		_user(DRAFTSMAN, "Dci Draftsman", ["Projects User", "Draftsman"])
		_user(PEC_ADMIN, "Dci Admin", ["Projects User", "PEC Administrator"])

	def tearDown(self):
		frappe.set_user("Administrator")
		super().tearDown()

	# helpers -----------------------------------------------------------------------------------

	def new_project(self, **values):
		return frappe.get_doc(
			{
				"doctype": "Project",
				"project_name": f"DCI test {frappe.generate_hash(length=8)}",
				"company": self.company,
				**values,
			}
		).insert(ignore_permissions=True)

	def new_scope(self, project, scope_name="Piping", **values):
		return frappe.get_doc(
			{"doctype": "Scope", "scope_name": scope_name, "project": project.name, **values}
		).insert(ignore_permissions=True)

	def new_task(self, project, subject="Piping layout", **values):
		return frappe.get_doc(
			{"doctype": "Task", "subject": subject, "project": project.name, **values}
		).insert(ignore_permissions=True)

	def submit(self, task, days_ago=0, **row):
		"""Add the next revision to `task` and save it."""
		task.append("pec_revisions", {"submission_date": add_days(today(), -days_ago), **row})
		return task.save(ignore_permissions=True)

	def answer(self, task, code, **row):
		"""Record the client's answer on the latest revision."""
		task.pec_revisions[-1].update({"approval_code": code, "received_date": today(), **row})
		return task.save(ignore_permissions=True)

	def report_row(self, project, task, **filters):
		rows = get_dci_data({"project": project.name, **filters}).rows
		return next(row for row in rows if row.task == task.name)

	# Approval Codes on Project ------------------------------------------------------------------

	def test_new_project_gets_default_codes(self):
		project = self.new_project()
		self.assertEqual([row.code for row in project.pec_approval_codes], ["1", "2", "3", "4"])
		self.assertEqual([row.code for row in project.pec_approval_codes if row.closes_document], ["1"])
		self.assertEqual(project.pec_approval_codes[2].description, "Revise and resubmit")

	def test_project_from_won_opportunity_gets_default_codes(self):
		customer = frappe.get_doc(
			{"doctype": "Customer", "customer_name": "DCI Test Customer", "customer_type": "Company"}
		).insert(ignore_permissions=True)
		opp = frappe.get_doc(
			{
				"doctype": "Opportunity",
				"opportunity_from": "Customer",
				"party_name": customer.name,
				"company": self.company,
				"title": f"DCI handoff {frappe.generate_hash(length=6)}",
			}
		).insert(ignore_permissions=True)

		project = frappe.get_doc("Project", create_project(opp))
		self.assertEqual([row.code for row in project.pec_approval_codes], ["1", "2", "3", "4"])

	def test_own_codes_are_kept_and_editable(self):
		codes = [{"code": "A", "closes_document": 1}, {"code": "B"}, {"code": "C"}]
		project = self.new_project(pec_approval_codes=codes)
		self.assertEqual([row.code for row in project.pec_approval_codes], ["A", "B", "C"])

		# defaults replaced after creation
		project = self.new_project()
		project.set("pec_approval_codes", codes)
		project.save(ignore_permissions=True)
		self.assertEqual([row.code for row in get_approval_codes(project.name)], ["A", "B", "C"])

	def test_duplicate_code_rejected(self):
		project = self.new_project()
		project.append("pec_approval_codes", {"code": " 2 ", "description": "again"})
		self.assertRaisesRegex(frappe.ValidationError, "already in the list", project.save)

	def test_code_in_use_cannot_be_removed(self):
		project = self.new_project()
		task = self.answer(self.submit(self.new_task(project)), "3")

		project.reload()
		project.set("pec_approval_codes", [row for row in project.pec_approval_codes if row.code != "3"])
		self.assertRaisesRegex(frappe.ValidationError, "cannot be removed", project.save)

		# an unused code can go
		project.reload()
		project.set("pec_approval_codes", [row for row in project.pec_approval_codes if row.code != "4"])
		project.save(ignore_permissions=True)
		self.assertEqual(task.reload().pec_latest_code, "3")

	def test_closing_flag_change_updates_task_status(self):
		project = self.new_project()
		task = self.answer(self.submit(self.new_task(project)), "2")
		self.assertEqual(task.pec_latest_status, "R0 Code 2 Received")

		project.reload()
		project.pec_approval_codes[1].closes_document = 1
		project.save(ignore_permissions=True)
		self.assertEqual(task.reload().pec_latest_status, "Approved (R0, Code 2)")
		self.assertEqual(self.report_row(project, task).latest_status, "Approved (R0, Code 2)")

	# revisions on Task --------------------------------------------------------------------------

	def test_latest_status_follows_the_revisions(self):
		project = self.new_project()
		task = self.new_task(project)
		self.assertEqual(task.pec_latest_status, "Not Submitted")
		self.assertIsNone(task.pec_latest_revision)

		self.submit(task, days_ago=10)
		self.assertEqual(task.pec_revisions[0].revision, "R0")
		self.assertEqual(task.pec_latest_status, "R0 Submitted – Awaiting Comments")
		self.assertEqual(self.report_row(project, task).latest_status, "R0 Submitted – Awaiting Comments")

		self.answer(task, "3")
		self.assertEqual(task.pec_latest_status, "R0 Code 3 Received")
		self.assertEqual(task.pec_revisions[0].code_description, "Revise and resubmit")

		self.answer(self.submit(task), "1")
		self.assertEqual(task.pec_latest_status, "Approved (R1, Code 1)")
		self.assertEqual(
			(task.pec_latest_revision, task.pec_latest_code, getdate(task.pec_latest_received_date)),
			("R1", "1", getdate(today())),
		)

	def test_revisions_do_not_change_status_or_progress(self):
		task = self.new_task(self.new_project(), status="Working", progress=40)
		self.answer(self.submit(task), "1")
		self.assertEqual((task.status, task.progress), ("Working", 40))

	def test_next_revision_needs_a_code_on_the_previous_one(self):
		task = self.submit(self.new_task(self.new_project()))
		task.append("pec_revisions", {"submission_date": today()})
		self.assertRaisesRegex(frappe.ValidationError, "only after", task.save)

		task.reload()
		self.answer(task, "3")
		self.submit(task)
		self.assertEqual([row.revision for row in task.pec_revisions], ["R0", "R1"])

	def test_earlier_revision_is_locked(self):
		task = self.submit(self.answer(self.submit(self.new_task(self.new_project())), "3"))

		frappe.set_user(ENGINEER)
		task = frappe.get_doc("Task", task.name)
		task.pec_revisions[0].notes = "changed afterwards"
		self.assertRaisesRegex(frappe.ValidationError, "R0.* is locked", task.save)

		task.reload()
		task.remove(task.pec_revisions[0])
		self.assertRaisesRegex(frappe.ValidationError, "is locked", task.save)

		# the latest revision stays open, and so does the rest of the Task
		task.reload()
		task.pec_revisions[1].notes = "sent by email"
		task.progress = 60
		task.save()

		frappe.set_user(PEC_ADMIN)
		task = frappe.get_doc("Task", task.name)
		task.pec_revisions[0].notes = "corrected by the administrator"
		task.save()
		self.assertEqual(task.pec_revisions[0].notes, "corrected by the administrator")

	def test_at_most_eleven_revisions(self):
		task = self.new_task(self.new_project())
		for _n in range(11):
			task.append("pec_revisions", {"submission_date": today(), "approval_code": "3"})
		task.save(ignore_permissions=True)
		self.assertEqual(task.pec_revisions[-1].revision, "R10")
		self.assertEqual(task.pec_latest_status, "R10 Code 3 Received")

		task.append("pec_revisions", {"submission_date": today()})
		self.assertRaisesRegex(frappe.ValidationError, "at most 11", task.save)

	def test_row_values_are_checked(self):
		project = self.new_project(pec_approval_codes=[{"code": "A", "closes_document": 1}, {"code": "B"}])
		task = self.submit(self.new_task(project))

		task.pec_revisions[0].approval_code = "1"
		self.assertRaisesRegex(frappe.ValidationError, "not one of Project", task.save)

		task.reload()
		task.pec_revisions[0].update({"approval_code": "B", "received_date": add_days(today(), -3)})
		self.assertRaisesRegex(frappe.ValidationError, "cannot be before", task.save)

		task.reload()
		self.answer(task, "A")
		self.assertEqual(task.pec_latest_status, "Approved (R0, Code A)")

		task.append("pec_revisions", {"approval_code": "B"})  # no Submission Date
		self.assertRaises(frappe.MandatoryError, task.save)

	def test_document_number_is_unique_in_a_project(self):
		project = self.new_project()
		self.new_task(project, pec_document_number="ABC123-PIP-001")
		with self.assertRaisesRegex(frappe.ValidationError, "already used"):
			self.new_task(project, pec_document_number=" ABC123-PIP-001 ")
		# another Project may use the same number
		self.new_task(self.new_project(), pec_document_number="ABC123-PIP-001")

	def test_document_number_is_searchable(self):
		meta = frappe.get_meta("Task")
		self.assertIn("pec_document_number", meta.get_search_fields())
		self.assertTrue(meta.get_field("pec_document_number").in_list_view)

	# DCI report ---------------------------------------------------------------------------------

	def dci_project(self):
		"""Two Scopes, four documents: approved at R1, code 3 on R0, awaiting comments, not submitted."""
		project = self.new_project(project_manager=MANAGER, end_user="Coastal Terminal Operator")
		frappe.set_user(MANAGER)  # the Scope's owner is the report's Manager
		piping = self.new_scope(project, "Piping", assigned_engineer=ENGINEER, associate=DRAFTSMAN)
		civil = self.new_scope(project, "Civil")
		frappe.set_user("Administrator")

		approved = self.new_task(project, "Layout", scope=piping.name, pec_document_number="P-PIP-002")
		self.answer(self.submit(self.answer(self.submit(approved, days_ago=9), "3"), days_ago=2), "1")
		commented = self.new_task(project, "Isometric", scope=piping.name, pec_document_number="P-PIP-001")
		self.answer(self.submit(commented, days_ago=5), "3")
		awaiting = self.new_task(project, "Foundation", scope=civil.name, pec_document_number="P-CIV-001")
		self.submit(awaiting, days_ago=1)
		not_submitted = self.new_task(project, "Unnumbered", scope=piping.name)

		frappe.get_doc(
			{
				"doctype": "ToDo",
				"allocated_to": PEC_ADMIN,
				"reference_type": "Task",
				"reference_name": approved.name,
				"description": "Draft it",
			}
		).insert(ignore_permissions=True)
		return frappe._dict(
			project=project,
			piping=piping,
			civil=civil,
			approved=approved,
			commented=commented,
			awaiting=awaiting,
			not_submitted=not_submitted,
		)

	def test_report_columns(self):
		columns, *_rest = execute({})
		fieldnames = [column["fieldname"] for column in columns]
		self.assertEqual(len(fieldnames), 7 + 33 + 5)
		self.assertEqual(len(set(fieldnames)), len(fieldnames))
		self.assertEqual(
			fieldnames[:10],
			[
				"sr_no",
				"scope_name",
				"document_number",
				"document_name",
				"manager",
				"engineer",
				"team",
				"r0_submission",
				"r0_received",
				"r0_code",
			],
		)
		self.assertEqual(fieldnames[37:40], ["r10_submission", "r10_received", "r10_code"])
		self.assertEqual(fieldnames[-1], "latest_status")
		self.assertEqual(columns[7]["label"], "R0 Submission Date")

	def test_report_rows_summary_and_chart(self):
		dci = self.dci_project()
		_columns, rows, _message, chart, summary = execute({"project": dci.project.name})

		# by Scope, then Document Number; the unnumbered document last in its Scope
		self.assertEqual(
			[(row.sr_no, row.document_number) for row in rows],
			[(1, "P-PIP-001"), (2, "P-PIP-002"), (3, None), (4, "P-CIV-001")],
		)
		row = rows[1]
		self.assertEqual((row.scope_name, row.document_name), ("Piping", "Layout"))
		self.assertEqual((row.manager, row.engineer), ("Dci Manager", "Dci Engineer"))
		self.assertEqual(row.team, "Dci Admin, Dci Draftsman")
		self.assertEqual((row.r0_code, row.r1_code), ("3", "1"))
		self.assertEqual(getdate(row.r0_submission), getdate(add_days(today(), -9)))
		self.assertEqual(getdate(row.r1_received), getdate(today()))
		self.assertNotIn("r2_submission", row)
		self.assertEqual((row.latest_revision, row.latest_code), ("R1", "1"))
		self.assertEqual(row.latest_status, "Approved (R1, Code 1)")
		self.assertEqual(rows[2].latest_status, "Not Submitted")

		self.assertEqual(
			{card["label"]: card["value"] for card in summary},
			{
				"Code 1": 1,
				"Code 2": 0,
				"Code 3": 1,
				"Code 4": 0,
				"Not Submitted": 1,
				"Awaiting Comments": 1,
				"Total Documents": 4,
			},
		)
		self.assertEqual(chart["type"], "bar")
		self.assertEqual(chart["data"]["labels"], ["Code 1", "Code 2", "Code 3", "Code 4"])
		self.assertEqual(chart["data"]["datasets"][0]["values"], [1, 0, 1, 0])

	def test_report_filters(self):
		dci = self.dci_project()

		def documents(**filters):
			return [row.document_name for row in get_dci_data({"project": dci.project.name, **filters}).rows]

		self.assertEqual(documents(scope=dci.civil.name), ["Foundation"])
		self.assertEqual(documents(engineer=ENGINEER), ["Isometric", "Layout", "Unnumbered"])
		self.assertEqual(documents(latest_code="3"), ["Isometric"])
		self.assertEqual(documents(latest_status="Approved"), ["Layout"])
		self.assertEqual(documents(latest_status="Awaiting Comments"), ["Foundation"])
		self.assertEqual(documents(latest_status="Not Submitted"), ["Unnumbered"])
		self.assertEqual(documents(latest_status="Code Received"), ["Isometric"])

	def test_report_reads_revisions_in_one_query(self):
		dci = self.dci_project()
		with self.assertQueryCount(7):
			get_dci_data({"project": dci.project.name})

	def test_report_lists_more_than_a_page_of_documents(self):
		project = self.new_project()
		for n in range(25):
			self.new_task(project, f"Document {n}", pec_document_number=f"D-{n:03d}")
		self.assertEqual(len(get_dci_data({"project": project.name}).rows), 25)

	def test_execution_team_restriction_applies_to_the_report(self):
		dci = self.dci_project()
		other = self.new_project()
		frappe.get_doc(
			{
				"doctype": "User Permission",
				"user": DRAFTSMAN,
				"allow": "Project",
				"for_value": other.name,
				"apply_to_all_doctypes": 1,
			}
		).insert(ignore_permissions=True)

		frappe.set_user(DRAFTSMAN)
		self.assertEqual(get_dci_data({"project": dci.project.name}).rows, [])
		with self.assertRaises(frappe.ValidationError):
			run_report(REPORT, filters={"project": dci.project.name}, js_filters=JS_FILTERS)

	# exports ------------------------------------------------------------------------------------

	def export(self, project, file_format_type):
		frappe.local.form_dict = frappe._dict(
			report_name=REPORT,
			file_format_type=file_format_type,
			filters=json.dumps({"project": project.name}),
		)
		export_query()
		return frappe.response.filecontent

	def test_standard_export_to_excel_and_csv(self):
		dci = self.dci_project()
		frappe.set_user(ENGINEER)

		sheet = load_workbook(BytesIO(self.export(dci.project, "Excel"))).active
		header = [cell.value for cell in sheet[1]]
		self.assertEqual(len(header), 45)
		self.assertEqual(header[:3], ["Sr. No.", "Scope Name", "Document Number"])
		self.assertEqual(header[7], "R0 Submission Date")
		self.assertEqual(sheet.max_row, 5)
		self.assertEqual(sheet.cell(row=3, column=3).value, "P-PIP-002")

		lines = self.export(dci.project, "CSV").decode().strip().splitlines()
		self.assertEqual(len(lines), 5)
		self.assertIn("P-PIP-002", lines[2])
		self.assertIn("Approved (R1, Code 1)", lines[2])

	def test_client_format_workbook(self):
		dci = self.dci_project()
		frappe.set_user(ENGINEER)
		download_dci_xlsx(dci.project.name)
		self.assertEqual(frappe.response.type, "binary")
		self.assertEqual(frappe.response.filename, f"DCI_{dci.project.name}_{today()}.xlsx")

		workbook = load_workbook(BytesIO(frappe.response.filecontent))
		self.assertEqual(workbook.sheetnames, ["DCI", "Code Summary"])
		sheet = workbook["DCI"]

		# project header block
		self.assertEqual(
			[(sheet.cell(row=n, column=1).value, sheet.cell(row=n, column=3).value) for n in range(1, 6)],
			[
				("Project Number", dci.project.name),
				("Project Name", dci.project.project_name),
				("Project Type", None),
				("Client Name", None),
				("End User Name", "Coastal Terminal Operator"),
			],
		)

		# row 7: one merged header per revision and for Latest Status; row 8: their columns
		merged = {str(cells) for cells in sheet.merged_cells.ranges}
		self.assertTrue({"A7:A8", "G7:G8", "H7:J7", "AL7:AN7", "AO7:AS7"} <= merged)
		self.assertEqual(
			(sheet["H7"].value, sheet["AL7"].value, sheet["AO7"].value), ("R0", "R10", "Latest Status")
		)
		self.assertEqual(
			[sheet.cell(row=8, column=n).value for n in range(8, 11)],
			["Submission Date", "Received Date", "Approval Code"],
		)
		self.assertEqual(sheet["AS8"].value, "Status")
		self.assertTrue(sheet["H7"].font.bold)
		self.assertEqual(sheet["H7"].border.left.style, "thin")
		self.assertEqual(sheet.freeze_panes, "H9")

		# one row per document, same order and values as the report
		self.assertEqual(sheet.max_row, 12)
		self.assertEqual([sheet.cell(row=10, column=n).value for n in (1, 3, 4)], [2, "P-PIP-002", "Layout"])
		self.assertEqual(sheet["H10"].value.date(), getdate(add_days(today(), -9)))
		self.assertEqual((sheet["J10"].value, sheet["M10"].value), ("3", "1"))
		self.assertEqual(sheet["AS10"].value, "Approved (R1, Code 1)")

		summary = [[cell.value for cell in row] for row in workbook["Code Summary"].iter_rows()]
		self.assertEqual(summary[0], ["Approval Code", "Description", "Documents"])
		self.assertEqual(summary[1], ["1", "No comments, proceed", 1])
		self.assertEqual(summary[-1], ["Total Documents", None, 4])

		# one Scope only
		sheet = load_workbook(BytesIO(build_dci_workbook(dci.project.name, dci.civil.name)))["DCI"]
		self.assertEqual((sheet["A6"].value, sheet["C6"].value), ("Scope", "Civil"))
		self.assertEqual(sheet.max_row, 9)

	def test_draftsman_can_view_but_not_export(self):
		dci = self.dci_project()
		frappe.set_user(DRAFTSMAN)

		result = run_report(REPORT, filters={"project": dci.project.name}, js_filters=JS_FILTERS)
		self.assertEqual(len(result["result"]), 4)
		self.assertTrue(result["report_summary"])

		self.assertFalse(frappe.permissions.can_export("Task"))
		self.assertRaises(frappe.PermissionError, self.export, dci.project, "Excel")
		self.assertRaises(frappe.PermissionError, self.export, dci.project, "CSV")
		self.assertRaises(frappe.PermissionError, download_dci_xlsx, dci.project.name)

		# updating a revision is progress updating: allowed
		task = frappe.get_doc("Task", dci.awaiting.name)
		task.pec_revisions[0].notes = "transmittal T-014"
		task.save()

	def test_task_permissions_of_the_pec_roles(self):
		def rights(role):
			perm = frappe.get_all(
				"Custom DocPerm",
				filters={"parent": "Task", "role": role, "permlevel": 0},
				fields=["read", "write", "report", "export"],
			)[0]
			return (perm.read, perm.write, perm.report, perm.export)

		self.assertEqual(rights("Projects Manager"), (1, 1, 1, 1))
		self.assertEqual(rights("Engineer"), (1, 1, 1, 1))
		self.assertEqual(rights("Draftsman"), (1, 1, 1, 0))
		self.assertEqual(rights("Associate"), (1, 0, 1, 0))
		# ERPNext's own rule is kept
		self.assertEqual(rights("Projects User"), (1, 1, 1, 0))

	# migration ----------------------------------------------------------------------------------

	def test_old_dci_revisions_are_copied_to_the_task(self):
		project = self.new_project()
		task = self.new_task(project)
		dci = frappe.get_doc(
			{
				"doctype": "DCI",
				"document_title": "Old register entry",
				"project": project.name,
				"task": task.name,
			}
		).insert(ignore_permissions=True)
		orphan = frappe.get_doc(
			{"doctype": "DCI", "document_title": "No task", "project": project.name}
		).insert(ignore_permissions=True)
		for revision_no, state in enumerate(["Revision Required", "Approved", "Under Review"]):
			frappe.get_doc(
				{
					"doctype": "PEC Revision",
					"naming_series": "REV-.YYYY.-.#####",
					"dci": dci.name,
					"revision_no": revision_no,
					"workflow_state": state,
					"submission_date": add_days(today(), revision_no - 10),
					"overall_comments": f"comment {revision_no}",
				}
			).db_insert()

		skipped = []
		self.assertEqual(dci_patch._copy_revisions(dci, skipped), project.name)
		self.assertIsNone(dci_patch._copy_revisions(orphan, skipped))
		self.assertIsNone(dci_patch._copy_revisions(dci, skipped))  # second run copies nothing
		refresh_latest_fields(project.name)

		task.reload()
		self.assertEqual([row.revision for row in task.pec_revisions], ["R0", "R1", "R2"])
		self.assertEqual([row.approval_code for row in task.pec_revisions], ["3", "1", None])
		self.assertEqual(getdate(task.pec_revisions[0].submission_date), getdate(add_days(today(), -10)))
		self.assertTrue(task.pec_revisions[1].received_date)
		self.assertIsNone(task.pec_revisions[2].received_date)
		self.assertIn("comment 0", task.pec_revisions[0].notes)
		self.assertEqual(task.pec_latest_status, "R2 Submitted – Awaiting Comments")
		self.assertEqual(len(skipped), 2)
		self.assertIn(orphan.name, skipped[0])

		# the old records are untouched, and the copied Task can be saved and continued
		self.assertEqual(frappe.db.count("PEC Revision", {"dci": dci.name}), 3)
		self.answer(task, "1")
		self.assertEqual(task.pec_latest_status, "Approved (R2, Code 1)")
