"""Client-format DCI workbook: the PEC DCI Report laid out like the client's DCI sheet, with a project
header block and one merged header per revision. It is built from the report's own data function, so
it always shows what the report shows. The plain export is the standard one (report Menu > Export).
"""

from io import BytesIO

import frappe
from frappe import _
from frappe.utils import today
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pew_customizations.pew_customizations.report.pec_dci_report.pec_dci_report import (
	LATEST_COLUMNS,
	LEAD_COLUMNS,
	REVISION_COLUMNS,
	get_counts,
	get_dci_data,
)
from pew_customizations.projects.task_revisions import MAX_REVISIONS

GROUP_HEADER_ROW = 7  # R0 ... R10 and Latest Status, each merged over its columns
COLUMN_HEADER_ROW = 8
DATE_FORMAT = "DD-MMM-YYYY"

_BOLD = Font(bold=True)
_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


@frappe.whitelist()
def download_dci_xlsx(project: str, scope: str | None = None):
	frappe.has_permission("Task", "export", throw=True)
	# same rule as the report's Project filter (frappe.desk.query_report.validate_filters_permissions)
	if not (
		frappe.has_permission("Project", "read", project)
		or frappe.has_permission("Project", "select", project)
	):
		frappe.throw(
			_("You do not have permission to access Project {0}.").format(project), frappe.PermissionError
		)

	frappe.response.filename = f"DCI_{project}_{today()}.xlsx"
	frappe.response.filecontent = build_dci_workbook(project, scope or None)
	frappe.response.type = "binary"


def build_dci_workbook(project, scope=None):
	"""The workbook as bytes. Kept separate from the download so other flows can attach the same
	file, e.g. when a Scope is assigned to an Associate / Sub-vendor by email."""
	if not frappe.db.exists("Project", project):
		frappe.throw(_("Project {0} not found").format(project), frappe.DoesNotExistError)
	dci = get_dci_data({"project": project, "scope": scope})

	workbook = Workbook()
	sheet = workbook.active
	sheet.title = "DCI"
	_write_project_header(sheet, project, scope)
	_write_table(sheet, dci.rows)
	_write_code_summary(workbook.create_sheet("Code Summary"), dci)

	output = BytesIO()
	workbook.save(output)
	return output.getvalue()


def _write_project_header(sheet, project, scope):
	details = frappe.db.get_value(
		"Project", project, ["project_name", "project_type", "customer", "end_user"], as_dict=True
	)
	client = frappe.db.get_value("Customer", details.customer, "customer_name") if details.customer else None
	header = [
		(_("Project Number"), project),
		(_("Project Name"), details.project_name),
		(_("Project Type"), details.project_type),
		(_("Client Name"), client),
		(_("End User Name"), details.end_user),
	]
	if scope:
		header.append((_("Scope"), frappe.db.get_value("Scope", scope, "scope_name")))

	# label over the first two columns, value over the next two
	for row, (label, value) in enumerate(header, start=1):
		sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
		sheet.merge_cells(start_row=row, start_column=3, end_row=row, end_column=4)
		sheet.cell(row=row, column=1, value=label).font = _BOLD
		sheet.cell(row=row, column=3, value=value)


def _write_table(sheet, rows):
	column = 1

	# leading columns: one header cell spanning both header rows
	for _fieldname, label, _fieldtype, width in LEAD_COLUMNS:
		_header(sheet, GROUP_HEADER_ROW, column, _(label), end_row=COLUMN_HEADER_ROW)
		sheet.column_dimensions[get_column_letter(column)].width = width / 7
		column += 1
	first_revision_column = column

	groups = [
		(f"R{n}", [(f"r{n}_{fieldname}", *rest) for fieldname, *rest in REVISION_COLUMNS])
		for n in range(MAX_REVISIONS)
	]
	groups.append((_("Latest Status"), LATEST_COLUMNS))

	fields = [(fieldname, fieldtype) for fieldname, _label, fieldtype, _width in LEAD_COLUMNS]
	for group_label, columns in groups:
		_header(sheet, GROUP_HEADER_ROW, column, group_label, end_column=column + len(columns) - 1)
		for fieldname, label, fieldtype, width in columns:
			# "Latest Submission Date" under the "Latest Status" group reads "Submission Date"
			_header(sheet, COLUMN_HEADER_ROW, column, _(label.removeprefix("Latest ")))
			sheet.column_dimensions[get_column_letter(column)].width = width / 7
			fields.append((fieldname, fieldtype))
			column += 1
	sheet.row_dimensions[COLUMN_HEADER_ROW].height = 30

	for row_number, row in enumerate(rows, start=COLUMN_HEADER_ROW + 1):
		for column, (fieldname, fieldtype) in enumerate(fields, start=1):
			cell = sheet.cell(row=row_number, column=column, value=row.get(fieldname))
			cell.border = _BORDER
			if fieldtype == "Date":
				cell.number_format = DATE_FORMAT

	# the headers and the document columns stay in view while scrolling through the revisions
	sheet.freeze_panes = sheet.cell(row=COLUMN_HEADER_ROW + 1, column=first_revision_column)


def _header(sheet, row, column, label, end_row=None, end_column=None):
	end_row, end_column = end_row or row, end_column or column
	if (end_row, end_column) != (row, column):
		sheet.merge_cells(start_row=row, start_column=column, end_row=end_row, end_column=end_column)
	sheet.cell(row=row, column=column, value=label)
	# every cell of a merged range carries the border and fill
	for cells in sheet.iter_rows(min_row=row, max_row=end_row, min_col=column, max_col=end_column):
		for cell in cells:
			cell.font = _BOLD
			cell.fill = _HEADER_FILL
			cell.border = _BORDER
			cell.alignment = _CENTER


def _write_code_summary(sheet, dci):
	counts = get_counts(dci)
	for column, label in enumerate([_("Approval Code"), _("Description"), _("Documents")], start=1):
		_header(sheet, 1, column, label)

	lines = [(code.code, code.description, count) for code, count in counts.codes]
	lines += [
		(_("Not Submitted"), None, counts.not_submitted),
		(_("Awaiting Comments"), None, counts.awaiting_comments),
		(_("Total Documents"), None, counts.total),
	]
	for row, line in enumerate(lines, start=2):
		for column, value in enumerate(line, start=1):
			sheet.cell(row=row, column=column, value=value).border = _BORDER
	for cell in sheet[len(lines) + 1]:
		cell.font = _BOLD

	for letter, width in (("A", 22), ("B", 45), ("C", 14)):
		sheet.column_dimensions[letter].width = width
